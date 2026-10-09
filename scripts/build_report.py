"""Authoring helper; uses the bundled Codex document/plot dependency runtime."""
from pathlib import Path
import json,re,os
ROOT=Path(__file__).resolve().parents[1]
os.environ.setdefault('MPLCONFIGDIR',str(ROOT/'artifacts/matplotlib'))
from reportlab.graphics.shapes import Drawing,Rect,String,Line,Polygon
from reportlab.graphics.charts.barcharts import HorizontalBarChart
from reportlab.graphics import renderSVG
from reportlab.graphics import renderPM
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import subprocess
ART=ROOT/'artifacts';ART.mkdir(exist_ok=True)
def save(drawing,name):
 svg=ART/(name+'.svg');svg.write_text(renderSVG.drawToString(drawing))
 renderPM.drawToFile(drawing, str(ART/(name+'.png')), fmt='PNG', dpi=240)
d=Drawing(650,355)
boxes=[(10,245,180,70,['Read only cloud APIs','AWS  Azure  Google']),(235,245,180,70,['Normalize observations','Source and timestamps']),(460,245,180,70,['Tri state policy checks','Passed  Failed  Unknown']),(460,125,180,70,['Explicit graph and score','Factors and uncertainty']),(235,125,180,70,['Snapshot store','SQLite and local audit']),(10,125,180,70,['Reviewer dashboard','Evidence and decisions']),(10,5,180,70,['What if simulation','Copied evidence only']),(235,5,180,70,['Bounded assistant','Validated finding IDs'])]
for x,y,w,h,labels in boxes:
 d.add(Rect(x,y,w,h,rx=5,ry=5,fillColor=colors.HexColor('#f0f4f7'),strokeColor=colors.HexColor('#9caebb')))
 for i,text in enumerate(labels):d.add(String(x+w/2,y+h/2+8-19*i,text,textAnchor='middle',fontName='Helvetica',fontSize=12,fillColor=colors.HexColor('#172735')))
def arrow(x1,y1,x2,y2):
 d.add(Line(x1,y1,x2,y2,strokeColor=colors.HexColor('#446474'),strokeWidth=1.5))
 if x2>x1:pts=[x2,y2,x2-7,y2+4,x2-7,y2-4]
 elif x2<x1:pts=[x2,y2,x2+7,y2+4,x2+7,y2-4]
 else:pts=[x2,y2,x2-4,y2+7,x2+4,y2+7]
 d.add(Polygon(pts,fillColor=colors.HexColor('#446474'),strokeColor=None))
for coords in [(190,280,235,280),(415,280,460,280),(550,245,550,195),(460,160,415,160),(235,160,190,160),(100,125,100,75),(325,125,325,75)]:arrow(*coords)
d.add(String(550,38,'No cloud write component',textAnchor='middle',fontSize=12,fillColor=colors.HexColor('#435b6a')));save(d,'architecture')
r=json.loads((ROOT/'evaluation/results.json').read_text());d=Drawing(650,245);chart=HorizontalBarChart();chart.x=145;chart.y=48;chart.width=445;chart.height=155
chart.data=[[m['ndcg_at_5'] for m in r['methods'].values()]];chart.categoryAxis.categoryNames=['Severity only','Context','Context and graph'];chart.categoryAxis.labels.fontSize=11
chart.valueAxis.valueMin=0;chart.valueAxis.valueMax=1;chart.valueAxis.valueStep=.25;chart.valueAxis.labels.fontSize=10
chart.bars[0].fillColor=colors.HexColor('#376358');chart.bars[0].strokeColor=None;chart.barLabelFormat='%.4f';chart.barLabels.fontSize=10;d.add(chart)
d.add(String(325,15,'NDCG at five on authored relevance labels',fontSize=12,textAnchor='middle'));save(d,'ranking')

def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr();x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);tcPr.append(x)
def keep_row(row):
 p=row._tr.get_or_add_trPr();p.append(OxmlElement('w:cantSplit'))
def hyperlink(p,label,url):
 rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'235E78');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def equation(doc,text):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 om=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);om.append(r);p._p.append(om)
 p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(9)
def table(doc,rows):
 t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=([1.1,5.4] if len(rows[0])==2 else [1.35,2.35,2.8] if len(rows[0])==3 else [1.15,1.9,1.85,1.6])
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 props=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for side in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 props.append(borders)
 margin=OxmlElement('w:tblCellMar')
 for side in ('top','bottom','left','right'):
  e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');margin.append(e)
 props.append(margin)
 for i,record in enumerate(rows):
  row=t.rows[0] if i==0 else t.add_row();keep_row(row)
  if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,(cell,text) in enumerate(zip(row.cells,record)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   shade(cell,'24465A' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'))
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
   run=p.add_run(text);run.font.name='Calibri';run.font.size=Pt(9.5);run.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '172735')
 doc.add_paragraph().paragraph_format.space_after=Pt(3)

doc=Document();doc.core_properties.title='Evidence Guided Multi Cloud Security Posture Assessment';doc.core_properties.author='Technical revision for candidate review';doc.core_properties.subject='Multi cloud capstone implementation and research';sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.9);sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.4)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
for n,size in [('Title',24),('Subtitle',13),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
 st=doc.styles[n];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
 if n.startswith('Heading'):st.font.bold=True;st.paragraph_format.keep_with_next=True
sec.different_first_page_header_footer=True
h=sec.header.paragraphs[0];h.text='Evidence Guided Multi Cloud Security Posture Assessment';h.style='Normal'
for r0 in h.runs:r0.font.size=Pt(8);r0.font.color.rgb=RGBColor(0,0,0)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r0=f.add_run('Candidate review edition   |   ');r0.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
lines=(ROOT/'docs/REVISED_REPORT.md').read_text().splitlines();i=0;first=True;pending_break=False
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line=='<!--page-->':pending_break=True;continue
 if line.startswith('|'):
  rows=[]
  while True:
   cells=[x.strip() for x in line.strip('|').split('|')]
   if not all(re.match(r'^:?-+:?$',x) for x in cells):rows.append(cells)
   if i>=len(lines) or not lines[i].strip().startswith('|'):break
   line=lines[i].strip();i+=1
  table(doc,rows);continue
 if line.startswith('[equation]'):equation(doc,line[len('[equation]'):]);continue
 if line.startswith('[figure:'):
  name=line[8:-1];p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(str(ART/(('dashboard-report' if name=='dashboard' else name)+'.png')),width=Inches(6.3));continue
 if line.startswith('# '):
  title=line[2:];p=doc.add_paragraph(title,'Title' if first else 'Heading 1');p.paragraph_format.page_break_before=pending_break;pending_break=False;first=False;continue
 if line.startswith('## '):doc.add_paragraph(line[3:],'Heading 2');continue
 p=doc.add_paragraph()
 url=re.search(r'https://\S+',line)
 if url:
  p.add_run(line[:url.start()]);hyperlink(p,'Source',url.group(0));p.add_run(line[url.end():])
 else:p.add_run(line)
 if line.startswith('Figure '):
  for r0 in p.runs:r0.font.size=Pt(9);r0.italic=True
 if line.startswith('[') and re.match(r'^\[\d+\]',line):p.paragraph_format.space_after=Pt(11)
# Title is plain black, no theme underline/border residue.

for style in doc.styles:
 for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 for fonts in style.element.xpath('.//w:rFonts'):
  for attribute in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
   if qn('w:'+attribute) in fonts.attrib:del fonts.attrib[qn('w:'+attribute)]
output=ROOT/'deliverables/Revised_Multi_Cloud_CSPM_Capstone.docx';doc.save(output)
print(f'Saved {output}; {len(doc.paragraphs)} paragraphs; {len(doc.tables)} tables')
